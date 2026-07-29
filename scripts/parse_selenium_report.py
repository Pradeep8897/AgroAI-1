from openpyxl import load_workbook
import sys

def main():
    path = 'selenium-tests/test-results/selenium-test-report.xlsx'
    try:
        wb = load_workbook(path)
    except Exception as e:
        print('ERROR loading workbook:', e)
        sys.exit(1)
    if 'Details' not in wb.sheetnames:
        print('Details sheet not found. Sheets:', wb.sheetnames)
        sys.exit(1)
    ws = wb['Details']
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        idx_id = headers.index('Test ID')
        idx_status = headers.index('Status')
    except ValueError:
        print('Could not find headers in Details sheet:', headers)
        sys.exit(1)
    fails = []
    passes = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        tid = row[idx_id]
        status = row[idx_status]
        if status == 'FAIL':
            fails.append(tid)
        elif status == 'PASS':
            passes += 1
    print('Passed', passes, 'Failed', len(fails))
    for t in fails[:500]:
        print(t)

if __name__ == '__main__':
    main()
