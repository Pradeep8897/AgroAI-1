import os
from openpyxl import load_workbook

base = os.path.join('artifacts','python-excel-reports')
if not os.path.isdir(base):
    print('No extracted folder at', base)
    raise SystemExit(0)

for root, dirs, files in os.walk(base):
    for f in files:
        if not f.lower().endswith('.xlsx'):
            continue
        path = os.path.join(root, f)
        print('\nFILE:', path)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            sheets = wb.sheetnames
            print('Sheets:', sheets)
            sheet = sheets[0]
            ws = wb[sheet]
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                row_count += 1
                cells = [str(c) if c is not None else '' for c in row]
                print(' | '.join(cells))
                if row_count >= 5:
                    break
        except Exception as e:
            print('Error reading', path, e)
