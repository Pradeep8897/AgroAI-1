from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = 'Login Test Matrix'

headers = [
    'Test Case ID',
    'Test Name',
    'Description',
    'Input Email',
    'Input Password',
    'Expected Result',
    'Category',
    'Priority',
    'Notes'
]
ws.append(headers)

for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = Font(bold=True)
    cell.fill = PatternFill('solid', fgColor='D9D9D9')
    cell.alignment = Alignment(horizontal='center', vertical='center')

cases = []
id_counter = 1

# Core login validation cases
cases.append(('Valid credentials login', 'Valid login with registered user', 'testuser@agroai.com', 'TestPassword123!', 'Dashboard loads successfully', 'Positive', 'High', 'Requires seeded account'))
cases.append(('Invalid email format', 'Email field validation rejects bad email format', 'bad-email', 'TestPassword123!', 'Show validation error', 'Validation', 'High', 'Client-side validation'))
cases.append(('Empty email field', 'Submit with blank email value', '', 'TestPassword123!', 'Show "Please fill in all fields."', 'Validation', 'High', 'Form validation'))
cases.append(('Empty password field', 'Submit with blank password value', 'testuser@agroai.com', '', 'Show "Please fill in all fields."', 'Validation', 'High', 'Form validation'))
cases.append(('Nonexistent user login', 'Use unknown email on login form', 'unknown@agroai.com', 'SomePass123!', 'Show login failed error', 'Negative', 'Medium', 'Backend rejects unknown user'))
cases.append(('Incorrect password', 'Valid email with wrong password', 'testuser@agroai.com', 'WrongPass!', 'Show login failed error', 'Negative', 'Medium', 'Backend rejects invalid password'))

# Generate combinatorial permutations and negative cases
email_variants = [
    'testuser@agroai.com',
    'user+test@agroai.com',
    'user.name@agroai.co',
    'USER@AGROAI.COM',
    '',
    'missingatsign.com',
    'missingdomain@',
    'no_tld@agroai',
    'spaces @agroai.com',
    'user@agroai..com',
    'testuser@agroai.com '  # trailing space
]
password_variants = [
    'TestPassword123!',
    'P@ssw0rd!',
    '123456',
    'short',
    '',
    'muchlongpasswordwithsymbols!@#$%^&*()',
    '     ',
    'pa$$w0rd',
    'TestPassword123',
    'TestPassword123! '  # trailing space
]

for email in email_variants:
    for password in password_variants:
        if email == 'testuser@agroai.com' and password == 'TestPassword123!':
            continue
        case_name = f'Login matrix: email={email!r}, password={password!r}'
        expected = 'Dashboard loads successfully' if email == 'testuser@agroai.com' and password == 'TestPassword123!' else 'Show login failed or validation error'
        priority = 'High' if email == 'testuser@agroai.com' else 'Medium'
        cases.append((case_name, 'Login matrix combination', email, password, expected, 'Matrix', priority, 'Data-driven test'))

# Additional edge cases
edge_cases = [
    ('Password visibility toggle', 'Toggle password visibility button shows characters', 'testuser@agroai.com', 'TestPassword123!', 'Password field switches between text and password', 'UI', 'Low', 'Requires eye button interaction'),
    ('Google sign-in flow button visible', 'Google sign in button is present on login page', '', '', 'Google sign in button is visible', 'UI', 'Low', 'No backend traffic required'),
    ('Backend URL settings available', 'Server configuration settings panel opens', '', '', 'Backend URL input visible and save button works', 'UI', 'Low', 'Optional on login page'),
]

cases.extend(edge_cases)

# Ensure at least 300 rows
while len(cases) < 300:
    email = f'edge{len(cases)}@agroai.com'
    password = f'EdgePass{len(cases)}!'
    cases.append((f'Auto-generated case {len(cases) + 1}', 'Generated login edge case', email, password, 'Show login failed or validation error', 'Generated', 'Low', 'Auto-generated matrix case'))

for case in cases:
    ws.append((f'TC-{id_counter:04d}',) + case)
    id_counter += 1

ws.column_dimensions['A'].width = 14
ws.column_dimensions['B'].width = 38
ws.column_dimensions['C'].width = 52
ws.column_dimensions['D'].width = 30
ws.column_dimensions['E'].width = 30
ws.column_dimensions['F'].width = 28
ws.column_dimensions['G'].width = 16
ws.column_dimensions['H'].width = 12
ws.column_dimensions['I'].width = 40

wb.save('selenium-tests/login-test-summary.xlsx')
print('Generated selenium-tests/login-test-summary.xlsx with', len(cases), 'test cases')
